from __future__ import annotations

import hashlib
import html
import json
import re
import uuid
from importlib import import_module
from pathlib import Path
from typing import Any

from eurus.coding.document_tools import MAX_DOCUMENT_BYTES, _read_document, _resolve
from eurus.errors import ToolError
from eurus.tools.base import Tool, ToolContext

MAX_ROWS = 10_000
MAX_COLUMNS = 100
MAX_SLIDES = 40


def _load_json(value: Any, label: str) -> Any:
    if isinstance(value, str):
        if len(value.encode("utf-8")) > MAX_DOCUMENT_BYTES:
            raise ToolError(f"{label} exceeds the 2 MB limit")
        try:
            return json.loads(value)
        except json.JSONDecodeError as exc:
            raise ToolError(f"{label} must be valid JSON: {exc.msg}") from exc
    if isinstance(value, (dict, list)):
        return value
    raise ToolError(f"{label} must be a JSON object or array")


def _target(context: ToolContext, raw_path: str, extension: str) -> Path:
    path = _resolve(context, raw_path)
    if path.suffix.lower() != extension:
        raise ToolError(f"path must end with {extension}")
    if path.exists():
        raise ToolError("file already exists; choose a new path to avoid overwriting")
    return path


def _save_guard(path: Path, save: Any) -> None:
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        save(path)
    except ToolError:
        raise
    except Exception as exc:
        raise ToolError(f"cannot create {path.suffix} file: {type(exc).__name__}: {exc}") from exc


def _document_class() -> Any:
    try:
        return import_module("docx").Document
    except ImportError as exc:
        raise ToolError("Word support requires `pip install 'eurus-agent[documents]'`") from exc


def _parse_word_content(content: str) -> list[tuple[str, str]]:
    parsed: list[tuple[str, str]] = []
    for line in content.splitlines():
        stripped = line.strip()
        if not stripped:
            parsed.append(("blank", ""))
        elif stripped.startswith("### "):
            parsed.append(("h3", stripped[4:]))
        elif stripped.startswith("## "):
            parsed.append(("h2", stripped[3:]))
        elif stripped.startswith("# "):
            parsed.append(("h1", stripped[2:]))
        elif stripped.startswith(("- ", "* ")):
            parsed.append(("bullet", stripped[2:]))
        elif re.match(r"^\d+[.)]\s+", stripped):
            parsed.append(("number", re.sub(r"^\d+[.)]\s+", "", stripped)))
        else:
            parsed.append(("paragraph", line))
    return parsed


def _render_word(spec: Any, path: Path) -> None:
    if isinstance(spec, str):
        spec = {"title": path.stem.replace("_", " ").title(), "content": spec}
    if not isinstance(spec, dict):
        raise ToolError("Word content must be a string or JSON object")
    title = str(spec.get("title", path.stem.replace("_", " ").title()))[:300]
    subtitle = str(spec.get("subtitle", ""))[:500]
    content = str(spec.get("content", ""))
    if not content.strip() and not spec.get("sections"):
        raise ToolError("Word content or sections must not be empty")
    if len(content.encode("utf-8")) > MAX_DOCUMENT_BYTES:
        raise ToolError("Word content exceeds the 2 MB limit")
    Document = _document_class()
    document = Document()
    section = document.sections[0]
    section.top_margin = import_module("docx.shared").Inches(0.7)
    section.bottom_margin = import_module("docx.shared").Inches(0.7)
    section.left_margin = import_module("docx.shared").Inches(0.8)
    section.right_margin = import_module("docx.shared").Inches(0.8)
    document.core_properties.title = title
    document.add_heading(title, 0)
    if subtitle:
        paragraph = document.add_paragraph(subtitle)
        paragraph.style = "Subtitle"
    for kind, text in _parse_word_content(content):
        if kind == "blank":
            continue
        if kind == "h1":
            document.add_heading(text, level=1)
        elif kind == "h2":
            document.add_heading(text, level=2)
        elif kind == "h3":
            document.add_heading(text, level=3)
        elif kind == "bullet":
            document.add_paragraph(text, style="List Bullet")
        elif kind == "number":
            document.add_paragraph(text, style="List Number")
        else:
            document.add_paragraph(text)
    sections = spec.get("sections", [])
    if not isinstance(sections, list):
        raise ToolError("Word sections must be an array")
    total_section_chars = 0
    for item in sections:
        if not isinstance(item, dict):
            raise ToolError("each Word section must be an object")
        heading = str(item.get("heading", ""))[:300]
        paragraphs = item.get("paragraphs", [])
        bullets = item.get("bullets", [])
        if not isinstance(paragraphs, list) or not isinstance(bullets, list):
            raise ToolError("Word section paragraphs and bullets must be arrays")
        total_section_chars += len(heading) + sum(len(str(value)) for value in paragraphs + bullets)
        if total_section_chars > MAX_DOCUMENT_BYTES:
            raise ToolError("Word sections exceed the 2 MB limit")
        if heading:
            document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(str(paragraph))
        for bullet in bullets:
            document.add_paragraph(str(bullet), style="List Bullet")
    document.save(path)


def _excel_class() -> Any:
    try:
        return import_module("openpyxl")
    except ImportError as exc:
        raise ToolError("Excel support requires `pip install 'eurus-agent[documents]'`") from exc


def _safe_cell(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    text = str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text[:32_000]


def _render_excel(spec: Any, path: Path) -> None:
    if not isinstance(spec, dict) or not isinstance(spec.get("sheets"), list):
        raise ToolError(
            "Excel data must be JSON like "
            '{"sheets":[{"name":"Overview","headers":[...],"rows":[...]}]}'
        )
    sheets = spec["sheets"]
    if not 1 <= len(sheets) <= 20:
        raise ToolError("Excel workbook must contain 1 to 20 sheets")
    openpyxl = _excel_class()
    styles = import_module("openpyxl.styles")
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    total_cells = 0
    for sheet_spec in sheets:
        if not isinstance(sheet_spec, dict):
            raise ToolError("each Excel sheet must be an object")
        name = str(sheet_spec.get("name", "Sheet"))[:31] or "Sheet"
        name = re.sub(r"[\\/*?:\[\]]", "_", name)
        base = name
        suffix = 2
        while name.casefold() in used:
            tail = f"_{suffix}"
            name = base[: 31 - len(tail)] + tail
            suffix += 1
        used.add(name.casefold())
        headers, rows = sheet_spec.get("headers", []), sheet_spec.get("rows", [])
        if not isinstance(headers, list) or not headers or len(headers) > MAX_COLUMNS:
            raise ToolError("each Excel sheet needs 1 to 100 headers")
        if not isinstance(rows, list) or len(rows) > MAX_ROWS:
            raise ToolError("each Excel sheet can contain at most 10000 data rows")
        total_cells += len(headers) * (len(rows) + 1)
        if total_cells > 100_000:
            raise ToolError("Excel workbook is limited to 100000 cells")
        worksheet = workbook.create_sheet(name)
        worksheet.append([_safe_cell(value) for value in headers])
        for cell in worksheet[1]:
            cell.font = styles.Font(bold=True, color="FFFFFF")
            cell.fill = styles.PatternFill("solid", fgColor="17324D")
            cell.alignment = styles.Alignment(vertical="center", wrap_text=True)
        worksheet.row_dimensions[1].height = 28
        for row in rows:
            if not isinstance(row, list) or len(row) > len(headers):
                raise ToolError("each Excel row must be an array no wider than its header")
            worksheet.append([_safe_cell(value) for value in row])
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column_cells in worksheet.columns:
            column_letter = column_cells[0].column_letter
            width = min(42, max(12, max(len(str(cell.value or "")) for cell in column_cells) + 2))
            worksheet.column_dimensions[column_letter].width = width
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = styles.Alignment(vertical="top", wrap_text=True)
    workbook.properties.title = str(spec.get("title", path.stem))[:250]
    workbook.save(path)


def _pptx_class() -> Any:
    try:
        return import_module("pptx")
    except ImportError as exc:
        raise ToolError(
            "PowerPoint support requires `pip install 'eurus-agent[documents]'`"
        ) from exc


def _pptx_text(
    slide: Any,
    pptx: Any,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    size: int,
    color: str,
    bold: bool = False,
) -> None:
    from pptx.util import Inches, Pt

    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = pptx.dml.color.RGBColor.from_string(color)


def _render_pptx(spec: Any, path: Path) -> list[dict[str, Any]]:
    if not isinstance(spec, dict) or not isinstance(spec.get("slides"), list):
        raise ToolError(
            "PowerPoint data must be JSON like "
            '{"title":"...","slides":[{"title":"...","bullets":[...]}]}'
        )
    slides: list[dict[str, Any]] = spec["slides"]
    if not 1 <= len(slides) <= MAX_SLIDES:
        raise ToolError("PowerPoint must contain 1 to 40 slides")
    pptx = _pptx_class()
    from pptx.util import Inches, Pt

    presentation = pptx.Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    palette = {
        "navy": "14283F",
        "paper": "F6F4EF",
        "teal": "167D83",
        "ink": "1C2938",
        "muted": "657386",
        "white": "FFFFFF",
    }
    title = str(spec.get("title", path.stem.replace("_", " ").title()))[:250]
    for index, slide_spec in enumerate(slides):
        if not isinstance(slide_spec, dict):
            raise ToolError("each PowerPoint slide must be an object")
        slide = presentation.slides.add_slide(blank)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = pptx.dml.color.RGBColor.from_string(palette["paper"])
        bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.18), Inches(7.5))
        bar.fill.solid()
        bar.fill.fore_color.rgb = pptx.dml.color.RGBColor.from_string(palette["teal"])
        bar.line.fill.background()
        heading = str(slide_spec.get("title", title if index == 0 else f"Key point {index}"))[:250]
        _pptx_text(
            slide,
            pptx,
            heading,
            0.75,
            0.55,
            11.8,
            0.75,
            size=30 if index == 0 else 26,
            color=palette["navy"],
            bold=True,
        )
        subtitle = str(slide_spec.get("subtitle", ""))[:400]
        if subtitle:
            _pptx_text(
                slide,
                pptx,
                subtitle,
                0.8,
                1.35,
                11.6,
                0.55,
                size=14,
                color=palette["muted"],
            )
        body_top = 2.05 if subtitle else 1.75
        body = str(slide_spec.get("body", ""))[:2500]
        if body:
            _pptx_text(slide, pptx, body, 0.85, body_top, 11.5, 1.15, size=17, color=palette["ink"])
            body_top += 1.25
        bullets = slide_spec.get("bullets", [])
        if not isinstance(bullets, list) or len(bullets) > 8:
            raise ToolError("each PowerPoint slide can contain at most 8 bullets")
        if bullets:
            box = slide.shapes.add_textbox(
                Inches(0.9),
                Inches(body_top),
                Inches(11.3),
                Inches(min(3.6, 0.48 * len(bullets) + 0.25)),
            )
            frame = box.text_frame
            frame.clear()
            frame.word_wrap = True
            for bullet_index, bullet in enumerate(bullets):
                paragraph = frame.paragraphs[0] if bullet_index == 0 else frame.add_paragraph()
                paragraph.text = f"•  {str(bullet)[:350]}"
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(18)
                paragraph.font.color.rgb = pptx.dml.color.RGBColor.from_string(palette["ink"])
                paragraph.space_after = Pt(12)
        callout = str(slide_spec.get("callout", ""))[:300]
        if callout:
            shape = slide.shapes.add_shape(1, Inches(0.85), Inches(6.1), Inches(11.6), Inches(0.72))
            shape.fill.solid()
            shape.fill.fore_color.rgb = pptx.dml.color.RGBColor.from_string(palette["navy"])
            shape.line.fill.background()
            shape.text_frame.text = callout
            paragraph = shape.text_frame.paragraphs[0]
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(15)
            paragraph.font.bold = True
            paragraph.font.color.rgb = pptx.dml.color.RGBColor.from_string(palette["white"])
        _pptx_text(
            slide,
            pptx,
            f"{index + 1:02d}  /  {len(slides):02d}",
            11.4,
            7.05,
            1.1,
            0.22,
            size=9,
            color=palette["muted"],
        )
    presentation.core_properties.title = title
    presentation.core_properties.subject = str(spec.get("subtitle", "Eurus presentation"))[:250]
    presentation.save(path)
    return slides


def _render_html_preview(spec: dict[str, Any], slides: list[dict[str, Any]], path: Path) -> None:
    title = html.escape(str(spec.get("title", path.stem))[:250])
    rendered: list[str] = []
    for index, item in enumerate(slides):
        heading = html.escape(str(item.get("title", f"Slide {index + 1}")))
        subtitle = html.escape(str(item.get("subtitle", "")))
        body = html.escape(str(item.get("body", "")))
        bullets = item.get("bullets", [])
        bullet_html = "".join(f"<li>{html.escape(str(value))}</li>" for value in bullets)
        callout = html.escape(str(item.get("callout", "")))
        rendered.append(
            f'<article class="slide"><canvas class="accent" width="20" height="720" aria-hidden="true"></canvas>'
            f'<div class="eyebrow">EURUS · FIELD NOTES · {index + 1:02d}</div><h2>{heading}</h2>'
            f'<p class="subtitle">{subtitle}</p><p>{body}</p><ul>{bullet_html}</ul>'
            f'<div class="callout">{callout}</div><span class="page">{index + 1:02d} / {len(slides):02d}</span></article>'
        )
    document = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{title} — Eurus presentation preview</title><style>
:root{{--ink:#1c2938;--navy:#14283f;--teal:#167d83;--paper:#f6f4ef;--muted:#657386}}
*{{box-sizing:border-box}}body{{margin:0;background:#e8e9e7;color:var(--ink);font:16px/1.55 Inter,ui-sans-serif,system-ui,sans-serif}}
header{{max-width:1100px;margin:32px auto 20px;padding:0 20px}}header p{{color:var(--muted)}}main{{display:grid;gap:24px;padding:0 20px 50px;max-width:1100px;margin:auto}}
.slide{{position:relative;overflow:hidden;min-height:500px;background:var(--paper);padding:60px 7% 55px;box-shadow:0 12px 35px #14283f18;border-radius:5px}}
.accent{{position:absolute;inset:0 auto 0 0;width:12px;height:100%;background:var(--teal)}}.eyebrow{{font-size:12px;letter-spacing:.15em;color:var(--teal);font-weight:700}}
h2{{font-size:clamp(30px,4vw,46px);line-height:1.08;color:var(--navy);max-width:850px;margin:24px 0 12px}}.subtitle{{color:var(--muted);font-size:18px}}
li{{margin:.6em 0}}.callout{{margin-top:24px;background:var(--navy);color:white;padding:14px 18px;min-height:16px}}.page{{position:absolute;right:28px;bottom:18px;color:var(--muted);font-size:11px}}
@media(max-width:600px){{.slide{{padding:38px 25px 50px;min-height:430px}}h2{{font-size:29px}}}}
</style></head><body><header><h1>{title}</h1><p>Interactive-ready HTML preview generated alongside the PowerPoint by Eurus.</p></header><main>{"".join(rendered)}</main>
<script>document.querySelectorAll('canvas.accent').forEach(c=>{{const x=c.getContext('2d');x.fillStyle='#167d83';x.fillRect(0,0,12,c.height);x.fillStyle='rgba(255,255,255,.3)';for(let y=24;y<c.height;y+=48)x.fillRect(12,y,8,1)}});</script>
</body></html>"""
    path.write_text(document, encoding="utf-8")


def create_office_artifact(
    context: ToolContext, format_name: str, path_text: str, payload: Any
) -> str:
    extensions = {"word": ".docx", "excel": ".xlsx", "powerpoint": ".pptx"}
    if format_name not in extensions:
        raise ToolError("format must be word, excel, or powerpoint")
    path = _target(context, path_text, extensions[format_name])
    if format_name == "word" and isinstance(payload, str):
        if payload.lstrip().startswith("{"):
            try:
                spec = json.loads(payload)
            except json.JSONDecodeError:
                spec = payload
        else:
            spec = payload
    else:
        spec = _load_json(payload, f"{format_name} content")
    if format_name == "word":
        _save_guard(path, lambda target: _render_word(spec, target))
        return f"created {path.relative_to(context.workspace.resolve()).as_posix()}"
    if format_name == "excel":
        _save_guard(path, lambda target: _render_excel(spec, target))
        return f"created {path.relative_to(context.workspace.resolve()).as_posix()}"
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        slides = _render_pptx(spec, path)
    except ToolError:
        raise
    except Exception as exc:
        raise ToolError(f"cannot create PowerPoint file: {type(exc).__name__}: {exc}") from exc
    preview_path = path.with_suffix(".preview.html")
    if preview_path.exists():
        path.unlink(missing_ok=True)
        raise ToolError("HTML preview already exists; choose a different PowerPoint path")
    try:
        _render_html_preview(spec, slides, preview_path)
    except OSError as exc:
        path.unlink(missing_ok=True)
        raise ToolError(f"PowerPoint created but HTML preview failed: {exc}") from exc
    relative = path.relative_to(context.workspace.resolve()).as_posix()
    preview_relative = preview_path.relative_to(context.workspace.resolve()).as_posix()
    return f"created {relative} and HTML Canvas design preview {preview_relative}"


async def _office_replace(arguments: dict[str, Any], context: ToolContext) -> str:
    format_name = str(arguments["format"]).strip().lower()
    path = _resolve(context, str(arguments["path"]))
    extensions = {"word": ".docx", "excel": ".xlsx", "powerpoint": ".pptx"}
    if format_name not in extensions or path.suffix.lower() != extensions.get(format_name):
        raise ToolError(
            "format must match the file extension: word/.docx, excel/.xlsx, powerpoint/.pptx"
        )
    current = _read_document(path)
    actual = hashlib.sha256(current.encode("utf-8")).hexdigest()
    expected = str(arguments["base_sha256"]).strip().lower()
    if len(expected) != 64 or expected != actual:
        raise ToolError("base_sha256 mismatch; read the latest document and retry")
    payload_key = {
        "word": "content",
        "excel": "data_json",
        "powerpoint": "slides_json",
    }[format_name]
    if not arguments.get(payload_key):
        raise ToolError(f"{payload_key} is required for {format_name} revisions")
    preview = path.with_suffix(".preview.html")
    if format_name == "powerpoint" and preview.is_symlink():
        raise ToolError("symbolic-link previews are not allowed")
    temp_name = f".{path.stem}.eurus-{uuid.uuid4().hex}{path.suffix}"
    temp_path = path.parent / temp_name
    create_office_artifact(
        context,
        format_name,
        temp_path.relative_to(context.workspace.resolve()).as_posix(),
        arguments[payload_key],
    )
    latest = _read_document(path)
    if hashlib.sha256(latest.encode("utf-8")).hexdigest() != expected:
        temp_path.unlink(missing_ok=True)
        temp_path.with_suffix(".preview.html").unlink(missing_ok=True)
        raise ToolError("document changed during revision; no replacement was applied")
    try:
        temp_path.replace(path)
        if format_name == "powerpoint":
            temp_preview = temp_path.with_suffix(".preview.html")
            if temp_preview.exists():
                temp_preview.replace(preview)
    except OSError as exc:
        raise ToolError(f"cannot replace Office document: {exc}") from exc
    relative = path.relative_to(context.workspace.resolve()).as_posix()
    if format_name == "powerpoint":
        preview_relative = preview.relative_to(context.workspace.resolve()).as_posix()
        return f"updated {relative} and {preview_relative} after SHA-256 verification"
    return f"updated {relative} after SHA-256 verification"


async def _word_create(arguments: dict[str, Any], context: ToolContext) -> str:
    return create_office_artifact(context, "word", str(arguments["path"]), arguments["content"])


async def _excel_create(arguments: dict[str, Any], context: ToolContext) -> str:
    return create_office_artifact(context, "excel", str(arguments["path"]), arguments["data_json"])


async def _powerpoint_create(arguments: dict[str, Any], context: ToolContext) -> str:
    return create_office_artifact(
        context, "powerpoint", str(arguments["path"]), arguments["slides_json"]
    )


def office_tools() -> list[Tool]:
    return [
        Tool(
            "word_create",
            (
                "Create a polished Word .docx file. Content can be Markdown-like text with "
                "headings and lists, or a JSON object with title, subtitle, content, and sections. "
                "Requires approval."
            ),
            {
                "type": "object",
                "properties": {"path": {"type": "string"}, "content": {"type": "string"}},
                "required": ["path", "content"],
                "additionalProperties": False,
            },
            _word_create,
            mutates=True,
            requires_approval=True,
        ),
        Tool(
            "excel_create",
            (
                "Create a formatted Excel .xlsx workbook from JSON: "
                "{title, sheets:[{name,headers,rows}]}. Includes styled headers, filters, frozen "
                "panes, sensible widths, and formula-injection protection. Requires approval."
            ),
            {
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "data_json": {"type": "string"},
                },
                "required": ["path", "data_json"],
                "additionalProperties": False,
            },
            _excel_create,
            mutates=True,
            requires_approval=True,
        ),
        Tool(
            "powerpoint_create",
            (
                "Create a designed PowerPoint .pptx from JSON "
                "{title,subtitle,slides:[{title,subtitle,body,bullets,callout}]}. Also creates a "
                "standalone responsive HTML preview with a Canvas accent. Requires approval."
            ),
            {
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "slides_json": {"type": "string"},
                },
                "required": ["path", "slides_json"],
                "additionalProperties": False,
            },
            _powerpoint_create,
            mutates=True,
            requires_approval=True,
        ),
        Tool(
            "office_replace",
            (
                "Revise an existing Word, Excel, or PowerPoint file only when base_sha256 matches "
                "the latest document_read result. Rebuilds the Office file from structured content "
                "and requires approval."
            ),
            {
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "format": {
                        "type": "string",
                        "enum": ["word", "excel", "powerpoint"],
                    },
                    "base_sha256": {"type": "string"},
                    "content": {"type": "string"},
                    "data_json": {"type": "string"},
                    "slides_json": {"type": "string"},
                },
                "required": ["path", "format", "base_sha256"],
                "additionalProperties": False,
            },
            _office_replace,
            mutates=True,
            requires_approval=True,
        ),
    ]
