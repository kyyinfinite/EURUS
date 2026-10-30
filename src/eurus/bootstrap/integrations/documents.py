from __future__ import annotations

import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from eurus.bootstrap.models import BROKEN, MISSING, OPTIONAL, READY, UNKNOWN, CheckResult


@dataclass(slots=True)
class DocumentsReport:
    docx: str = UNKNOWN
    openpyxl: str = UNKNOWN
    pptx: str = UNKNOWN
    creation_test: str = UNKNOWN
    detail: str = ""

    def to_dict(self) -> dict[str, Any]:
        return {
            "docx": self.docx,
            "openpyxl": self.openpyxl,
            "pptx": self.pptx,
            "creation_test": self.creation_test,
            "detail": self.detail,
        }


def _dist_state(dist_name: str) -> str:
    import importlib.metadata

    try:
        importlib.metadata.version(dist_name)
    except importlib.metadata.PackageNotFoundError:
        return MISSING
    return READY


def inspect_documents() -> DocumentsReport:
    report = DocumentsReport(
        docx=_dist_state("python-docx"),
        openpyxl=_dist_state("openpyxl"),
        pptx=_dist_state("python-pptx"),
    )
    if MISSING in {report.docx, report.openpyxl, report.pptx}:
        report.creation_test = UNKNOWN
        report.detail = "install the documents extra to enable office tools"
        return report
    ok, message = _creation_test()
    report.creation_test = READY if ok else BROKEN
    report.detail = message
    return report


def _creation_test() -> tuple[bool, str]:
    try:
        with tempfile.TemporaryDirectory(prefix="eurus-docs-") as folder:
            root = Path(folder)
            from docx import Document

            document = Document()
            document.add_paragraph("eurus bootstrap documents test")
            document.save(root / "test.docx")
            from openpyxl import Workbook

            workbook = Workbook()
            workbook.active["A1"] = "eurus"
            workbook.save(root / "test.xlsx")
            from pptx import Presentation

            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            box = slide.shapes.add_textbox(0, 0, 100, 100)
            box.text_frame.text = "eurus"
            presentation.save(root / "test.pptx")
            created = all(
                (root / name).is_file() for name in ("test.docx", "test.xlsx", "test.pptx")
            )
            if not created:
                return False, "test files were not created"
            return True, "DOCX, XLSX and PPTX created successfully"
    except Exception as exc:
        return False, f"creation test failed: {type(exc).__name__}: {exc}"


def documents_check_result(report: DocumentsReport) -> CheckResult:
    if MISSING in {report.docx, report.openpyxl, report.pptx}:
        return CheckResult(
            name="documents",
            status=OPTIONAL,
            detail="office document tools are not installed",
            hint="pip install -e '.[documents]'",
            optional=True,
        )
    if report.creation_test == BROKEN:
        return CheckResult(
            name="documents",
            status=BROKEN,
            detail=report.detail,
            optional=True,
        )
    return CheckResult(name="documents", status=READY, detail=report.detail, optional=True)
